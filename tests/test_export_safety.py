"""H2 spreadsheet-safe export and the I12 duplicate-column guard."""
import os
from types import SimpleNamespace as NS

import openpyxl
import pandas as pd
import pytest

from services.export_service import ExportService, spreadsheet_safe


def _frame():
    res = pd.DataFrame({
        "DEPTH": [1000.0, 1000.5, 1001.0],
        "GR": [-2.5, 50.0, 60.0],
        "ZONE": ["+cmd", "@SUM(A1)", "-5 ft"],
        "NOTE": ["=1+1", "-2.5", "\t=x"],
    })
    return res.rename(columns={"NOTE": "=BAD_CURVE"})


def _wells():
    res = _frame()

    def mk(name):
        return NS(display_name=name, key=name, results=res.copy(), calculated=True, status="run_ok",
                  summary={"gross_sand": 2.0, "net_pay": 1.0, "ng_pay": .5, "avg_phie_pay": .2,
                           "avg_sw_pay": .4, "hcpv_net_pay": .1, "rw": .05, "rw_source": "=1+1",
                           "rsh": 2, "rsh_source": "+cmd",
                           "zones": [{"zone": "@SUM(A1)", "top": 1, "bottom": 2,
                                      "params": {"rw": {"value": .05, "source": "=evil"}}}]},
                  qc_report=None, las_data=res[["DEPTH", "GR"]].copy(), well_info={"well_name": name})
    return [mk("=1+1"), mk("-WELL2")]


def _no_formulas(path):
    wb = openpyxl.load_workbook(path)
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", (ws.title, c.coordinate, c.value)
    return wb


def test_single_well_excel_has_no_formula_cells(tmp_path):
    svc = ExportService()
    errs = []
    svc.export_error.connect(errs.append)
    summary = {"well": "=1+1", "rw_source": "@SUM(A1)", "neg": -2.5, "txt": "-5 ft",
               "path": '=HYPERLINK("http://x")'}
    path = str(tmp_path / "single.xlsx")
    assert svc.export_excel(_frame(), summary, path), errs
    wb = _no_formulas(path)
    assert wb["Results"]["D1"].value == "=BAD_CURVE"
    assert wb["Results"]["D2"].value == "=1+1"
    assert wb["Results"]["B2"].value == -2.5
    values = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=2)}
    assert values["well"] == "=1+1"
    assert values["neg"] == -2.5


def test_multi_well_excel_has_no_formula_cells(tmp_path):
    svc = ExportService()
    errs = []
    svc.export_error.connect(errs.append)
    path = str(tmp_path / "multi.xlsx")
    assert svc.export_excel_wells(_wells(), path), errs
    wb = _no_formulas(path)
    assert wb["Summary"]["A2"].value == "=1+1"


def test_csv_prefix_rule(tmp_path):
    svc = ExportService()
    path = str(tmp_path / "r.csv")
    assert svc.export_csv(_frame(), path)
    back = pd.read_csv(path, dtype=str, keep_default_na=False)
    assert "'=BAD_CURVE" in back.columns
    assert list(back["ZONE"]) == ["'+cmd", "'@SUM(A1)", "'-5 ft"]
    assert back["'=BAD_CURVE"][0] == "'=1+1"
    assert back["'=BAD_CURVE"][1] == "-2.5"          # parses as a number: untouched
    assert back["GR"][0] == "-2.5"                   # numeric column untouched
    assert svc.get_csv_string(_frame()).splitlines()[0].endswith("'=BAD_CURVE")


def test_multi_well_csv_prefixes_well_names(tmp_path):
    svc = ExportService()
    path = str(tmp_path / "m.csv")
    assert svc.export_csv_wells(_wells(), path)
    back = pd.read_csv(path, dtype=str, keep_default_na=False)
    assert set(back["WELL"]) == {"'=1+1", "'-WELL2"}


def test_spreadsheet_safe_leaves_numeric_columns_and_input_alone():
    df = _frame()
    out = spreadsheet_safe(df, csv=True)
    assert out["GR"].tolist() == df["GR"].tolist()
    assert df["ZONE"][0] == "+cmd"                    # the original is not modified


@pytest.mark.parametrize("method", ["csv", "excel"])
def test_duplicate_columns_are_refused_with_the_names(tmp_path, method):
    svc = ExportService()
    errs = []
    svc.export_error.connect(errs.append)
    df = pd.DataFrame([[1, 2, 3]], columns=["DEPTH", "PHIE", "PHIE"])
    path = str(tmp_path / ("d.csv" if method == "csv" else "d.xlsx"))
    ok = svc.export_csv(df, path) if method == "csv" else svc.export_excel(df, {}, path)
    assert not ok and "PHIE" in errs[0] and "duplicate" in errs[0]
    assert not os.path.exists(path)
